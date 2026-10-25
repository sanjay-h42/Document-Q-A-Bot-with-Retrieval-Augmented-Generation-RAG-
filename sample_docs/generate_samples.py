import os
import pymupdf
from docx import Document

def create_sample_pdf(output_path: str):
    doc = pymupdf.open()
    
    # Page 1
    page1 = doc.new_page()
    page1.insert_text((50, 60), "Multimodal Artificial Intelligence: Foundations and Frontiers", fontsize=16, color=(0.1, 0.2, 0.6))
    page1.insert_text((50, 85), "Page 1: Introduction and Vision-Language Architectures", fontsize=12, color=(0.3, 0.3, 0.3))
    
    p1_content = (
        "1. Executive Overview:\n"
        "Multimodal AI refers to models capable of processing and relating information from multiple modalities,\n"
        "including text, computer vision, audio, and structured sensor streams. The unifying objective is to build\n"
        "cross-modal representations where semantic concepts are shared across modalities.\n\n"
        "2. Contrastive Pretraining:\n"
        "Pioneered by OpenAI's CLIP (Contrastive Language-Image Pre-Training), contrastive architectures train an\n"
        "image encoder (e.g., Vision Transformer ViT) and a text encoder simultaneously on web-scale pairs.\n"
        "The model maximizes the cosine similarity between matched (image, text) embeddings while minimizing similarity\n"
        "between mismatched negative pairs in the batch using a symmetric cross-entropy loss."
    )
    rect1 = pymupdf.Rect(50, 110, 550, 750)
    page1.insert_textbox(rect1, p1_content, fontsize=10)
    
    # Page 2
    page2 = doc.new_page()
    page2.insert_text((50, 60), "Multimodal Artificial Intelligence: Foundations and Frontiers", fontsize=16, color=(0.1, 0.2, 0.6))
    page2.insert_text((50, 85), "Page 2: Autoregressive Vision-Language Models and Cross-Attention", fontsize=12, color=(0.3, 0.3, 0.3))
    
    p2_content = (
        "3. Cross-Attention and Perceiver Resamplers:\n"
        "Modern Vision-Language Models (VLMs) such as Gemini, LLaVA, and Flamingo project visual patches\n"
        "into the token embedding space of large language models. Cross-attention layers enable the language\n"
        "decoder to attend directly to visual features during autoregressive token generation.\n\n"
        "4. Benchmark Evaluations:\n"
        "State-of-the-art multimodal systems are evaluated on rigorous benchmarks such as:\n"
        "- MMLU (Massive Multitask Language Understanding) for academic and reasoning performance.\n"
        "- MMMU (Massive Multi-discipline Multimodal Understanding) assessing university-level visual domain expertise.\n"
        "- MathVista for multimodal mathematical reasoning.\n\n"
        "5. Safety and Alignment:\n"
        "Multimodal models introduce novel attack vectors, such as typographic visual prompt injections, where adversarial\n"
        "text rendered inside images attempts to override safety boundaries. Robust alignment requires multimodal RLHF."
    )
    rect2 = pymupdf.Rect(50, 110, 550, 750)
    page2.insert_textbox(rect2, p2_content, fontsize=10)
    
    doc.save(output_path)
    doc.close()
    print(f"Generated PDF: {output_path}")

def create_sample_docx(output_path: str):
    doc = Document()
    doc.add_heading("Cloud-Native Architecture and High Availability Guide", level=0)
    
    doc.add_heading("1. Principles of Resilient Distributed Systems", level=1)
    doc.add_paragraph(
        "Modern enterprise cloud architectures require horizontal scalability, stateless services, and "
        "decoupled asynchronous communication. By deploying microservices across multiple availability zones (AZs), "
        "organizations minimize single points of failure and achieve 99.99% uptime SLAs."
    )
    
    doc.add_heading("2. Availability Tier Matrix", level=2)
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    
    headers = ["Service Tier", "Target Uptime (SLA)", "Recovery Time Objective (RTO)"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        
    data = [
        ["Tier 1 - Mission Critical", "99.99% (< 52.6 min downtime/year)", "< 5 minutes"],
        ["Tier 2 - Business Operational", "99.9% (< 8.76 hours downtime/year)", "< 1 hour"],
        ["Tier 3 - Internal Supporting", "99.0% (< 3.65 days downtime/year)", "< 24 hours"]
    ]
    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, cell_value in enumerate(row_data):
            table.rows[row_idx].cells[col_idx].text = cell_value
            
    doc.add_heading("3. Data Consistency Models and CAP Theorem", level=1)
    doc.add_paragraph(
        "According to Eric Brewer's CAP theorem, a distributed data store can only guarantee at most two out of "
        "Consistency, Availability, and Partition Tolerance simultaneously. Cloud distributed databases like Google "
        "Cloud Spanner leverage TrueTime atomic clocks to provide external consistency while maintaining high availability."
    )
    
    doc.save(output_path)
    print(f"Generated DOCX: {output_path}")

if __name__ == "__main__":
    out_dir = os.path.dirname(os.path.abspath(__file__))
    create_sample_pdf(os.path.join(out_dir, "multimodal_ai_overview.pdf"))
    create_sample_docx(os.path.join(out_dir, "cloud_architecture_guide.docx"))
