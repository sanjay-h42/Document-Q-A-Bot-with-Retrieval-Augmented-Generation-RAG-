import pytest
import io
try:
    import pymupdf as fitz
except ImportError:
    import fitz

from docx import Document as DocxDoc
from app.services.document_parser import DocumentParser

def test_parse_pdf():
    # Create sample 2-page PDF in memory
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 72), "First page: Artificial Intelligence and Machine Learning.")
    page2 = doc.new_page()
    page2.insert_text((50, 72), "Second page: Neural networks and Deep Learning architectures.")
    
    pdf_bytes = doc.tobytes()
    doc.close()
    
    parsed = DocumentParser.parse_pdf(pdf_bytes, "doc_pdf_1", "sample.pdf")
    assert parsed.document_id == "doc_pdf_1"
    assert parsed.filename == "sample.pdf"
    assert parsed.file_type == "pdf"
    assert parsed.total_pages == 2
    assert len(parsed.pages) == 2
    assert parsed.pages[0].page_number == 1
    assert "Artificial Intelligence" in parsed.pages[0].text
    assert parsed.pages[1].page_number == 2
    assert "Deep Learning" in parsed.pages[1].text

def test_parse_docx():
    docx_doc = DocxDoc()
    docx_doc.add_heading("Cloud Architecture Principles", level=1)
    docx_doc.add_paragraph("High availability and resilience are paramount.")
    table = docx_doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Tier"
    table.rows[0].cells[1].text = "SLA"
    table.rows[1].cells[0].text = "Production"
    table.rows[1].cells[1].text = "99.99%"
    
    buffer = io.BytesIO()
    docx_doc.save(buffer)
    docx_bytes = buffer.getvalue()
    
    parsed = DocumentParser.parse_docx(docx_bytes, "doc_docx_1", "arch.docx")
    assert parsed.document_id == "doc_docx_1"
    assert parsed.file_type == "docx"
    assert len(parsed.pages) >= 1
    assert "Cloud Architecture Principles" in parsed.pages[0].text
    assert "Production | 99.99%" in parsed.pages[0].text

def test_parse_txt():
    content = "Section 1: Quantum Computing Fundamentals.\x0cSection 2: Quantum Algorithms and Gates.".encode('utf-8')
    parsed = DocumentParser.parse_txt(content, "doc_txt_1", "quantum.txt")
    assert parsed.document_id == "doc_txt_1"
    assert parsed.total_pages == 2
    assert "Quantum Computing Fundamentals" in parsed.pages[0].text
    assert "Quantum Algorithms and Gates" in parsed.pages[1].text

def test_parse_unsupported():
    with pytest.raises(ValueError):
        DocumentParser.parse(b"dummy", "id_1", "bad.xyz", ".xyz")
